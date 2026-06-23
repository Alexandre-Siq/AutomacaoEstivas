from pathlib import Path

from openpyxl import load_workbook


BASE_DIR = Path(__file__).resolve().parent
MODELO_ORIGINAL = BASE_DIR / "Modelo_Solicitacao_SSHD.xlsx"
TEMPLATE_SAIDA = BASE_DIR / "TEMPLATE_NOVO.xlsx"
ABA_TEMPLATE = "SSHD"

CAMPOS_SOLICITANTE = {
    "B5": None,  # Nome do Solicitante
    "B6": None,  # SSHD
    "B7": None,  # Cargo
}

VALORES_FIXOS = {
    "B34": "COMPLEXO HOSPITALAR DOS ESTIVADORES",
    "B35": "SMS",
    "B37": "COMPLEXO HOSPITALAR DOS ESTIVADORES",
    "B38": "SMS",
}


def criar_template() -> Path:
    """Cria o template usado pela automacao sem alterar o layout original."""
    if not MODELO_ORIGINAL.exists():
        raise FileNotFoundError(f"Modelo original nao encontrado: {MODELO_ORIGINAL}")

    workbook = load_workbook(MODELO_ORIGINAL)

    if ABA_TEMPLATE not in workbook.sheetnames:
        raise ValueError(f"Aba obrigatoria nao encontrada: {ABA_TEMPLATE}")

    worksheet = workbook[ABA_TEMPLATE]

    for celula, valor in CAMPOS_SOLICITANTE.items():
        worksheet[celula] = valor

    for celula, valor in VALORES_FIXOS.items():
        worksheet[celula] = valor

    workbook.save(TEMPLATE_SAIDA)
    return TEMPLATE_SAIDA


if __name__ == "__main__":
    caminho_template = criar_template()
    print(f"Template criado em: {caminho_template}")
