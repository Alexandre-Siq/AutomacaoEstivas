from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .config import COLUNAS_FONTE_DADOS, VALORES_FIXOS_CAMPOS, VALORES_PADRAO
from .exceptions import ErroLeituraPlanilha
from .normalizacao import limpar_cabecalho, normalizar_texto, valor_excel, valor_vazio


LIMITE_BUSCA_CABECALHO = 50
CABECALHOS_OBRIGATORIOS = {
    "Nome Colaborador": {
        "colaborador",
        "funcionario",
        "nome",
        "nome_colaborador",
        "nome_completo",
        "nome_do_colaborador",
        "nome_do_funcionario",
        "nome_funcionario",
    },
    "CPF": {
        "cpf",
        "cpf_colaborador",
        "cpf_do_colaborador",
        "cpf_funcionario",
    },
    "Cargo": {
        "cargo",
        "cargo_funcao",
        "cargo_ou_funcao",
        "especialidade",
        "funcao",
    },
}


def _cabecalhos_da_linha(worksheet, numero_linha: int) -> set[str]:
    return {
        limpar_cabecalho(worksheet.cell(numero_linha, coluna).value)
        for coluna in range(1, worksheet.max_column + 1)
        if not valor_vazio(worksheet.cell(numero_linha, coluna).value)
    }


def _mapear_cabecalhos(worksheet, numero_linha: int) -> dict[str, list[int]]:
    cabecalhos: dict[str, list[int]] = {}

    for coluna in range(1, worksheet.max_column + 1):
        chave = limpar_cabecalho(worksheet.cell(numero_linha, coluna).value)
        if chave:
            cabecalhos.setdefault(chave, []).append(coluna)

    return cabecalhos


def _linha_tem_cabecalhos_obrigatorios(cabecalhos: set[str]) -> bool:
    return all(
        bool(cabecalhos.intersection(alias))
        for alias in CABECALHOS_OBRIGATORIOS.values()
    )


def _descrever_linhas_lidas(workbook) -> str:
    amostras: list[str] = []

    for worksheet in workbook.worksheets:
        for numero_linha in range(1, min(worksheet.max_row, LIMITE_BUSCA_CABECALHO) + 1):
            valores = [
                normalizar_texto(worksheet.cell(numero_linha, coluna).value)
                for coluna in range(1, min(worksheet.max_column, 12) + 1)
                if not valor_vazio(worksheet.cell(numero_linha, coluna).value)
            ]
            if valores:
                amostras.append(
                    f"Aba '{worksheet.title}', linha {numero_linha}: " + " | ".join(valores[:8])
                )
            if len(amostras) >= 8:
                return "\n".join(amostras)

    return "Nenhum texto foi encontrado nas primeiras linhas das abas."


def _encontrar_planilha_e_linha_cabecalho(workbook) -> tuple[Any, int]:
    for worksheet in workbook.worksheets:
        for numero_linha in range(1, min(worksheet.max_row, LIMITE_BUSCA_CABECALHO) + 1):
            cabecalhos = _cabecalhos_da_linha(worksheet, numero_linha)
            if _linha_tem_cabecalhos_obrigatorios(cabecalhos):
                return worksheet, numero_linha

    amostras = _descrever_linhas_lidas(workbook)
    raise ErroLeituraPlanilha(
        "Não foi possível localizar o cabeçalho da planilha fonte. "
        "Procurei por colunas equivalentes a Nome Colaborador, CPF e Cargo "
        f"nas primeiras {LIMITE_BUSCA_CABECALHO} linhas de todas as abas.\n\n"
        "Amostra do que foi encontrado:\n"
        f"{amostras}"
    )


def _primeira_coluna(
    mapa_cabecalhos: dict[str, list[int]],
    aliases: set[str],
    *,
    apos_coluna: int | None = None,
) -> int | None:
    candidatas: list[int] = []

    for alias in aliases:
        candidatas.extend(mapa_cabecalhos.get(alias, []))

    candidatas = sorted(set(candidatas))
    if apos_coluna is not None:
        posteriores = [coluna for coluna in candidatas if coluna > apos_coluna]
        if posteriores:
            return posteriores[0]

    return candidatas[0] if candidatas else None


def _montar_mapa_colunas(mapa_cabecalhos: dict[str, list[int]]) -> dict[str, int]:
    colunas: dict[str, int] = {}

    for campo, aliases in COLUNAS_FONTE_DADOS.items():
        if campo == "numero":
            continue

        coluna = _primeira_coluna(mapa_cabecalhos, aliases)
        if coluna is not None:
            colunas[campo] = coluna

    coluna_logradouro = colunas.get("logradouro")
    coluna_numero = _primeira_coluna(
        mapa_cabecalhos,
        COLUNAS_FONTE_DADOS["numero"],
        apos_coluna=coluna_logradouro,
    )
    if coluna_numero is not None:
        colunas["numero"] = coluna_numero

    return colunas


def _separar_logradouro(valor: Any) -> tuple[Any, Any, Any]:
    texto = normalizar_texto(valor)
    if not texto:
        return None, None, None

    padroes = [
        r"^(?P<logradouro>.+?)(?:,\s*)?(?:n[ºo°.]?|numero|número)\s*"
        r"(?P<numero>\d+[A-Za-z0-9/-]*)(?:\s*[-,]\s*(?P<complemento>.+))?$",
        r"^(?P<logradouro>.+?),\s*(?P<numero>\d+[A-Za-z0-9/-]*)"
        r"(?:\s*[-,]\s*(?P<complemento>.+))?$",
    ]

    for padrao in padroes:
        resultado = re.match(padrao, texto, flags=re.IGNORECASE)
        if resultado:
            return (
                normalizar_texto(resultado.group("logradouro")),
                normalizar_texto(resultado.group("numero")),
                normalizar_texto(resultado.group("complemento")),
            )

    return texto, None, None


def _ler_registro(row: int, worksheet, colunas: dict[str, int]) -> dict[str, Any]:
    registro = {
        campo: valor_excel(worksheet.cell(row, coluna).value)
        for campo, coluna in colunas.items()
    }

    logradouro, numero_extraido, complemento_extraido = _separar_logradouro(
        registro.get("logradouro")
    )
    if logradouro:
        registro["logradouro"] = logradouro
    if valor_vazio(registro.get("numero")) and numero_extraido:
        registro["numero"] = numero_extraido
    if valor_vazio(registro.get("complemento")) and complemento_extraido:
        registro["complemento"] = complemento_extraido

    registro.update(VALORES_FIXOS_CAMPOS)

    for campo, valor_padrao in VALORES_PADRAO.items():
        if valor_vazio(registro.get(campo)):
            registro[campo] = valor_padrao

    return registro


def ler_planilha_geral(caminho_planilha: str | Path) -> list[dict[str, Any]]:
    caminho = Path(caminho_planilha)

    if not caminho.exists():
        raise ErroLeituraPlanilha(f"Planilha não encontrada: {caminho}")
    if caminho.suffix.lower() != ".xlsx":
        raise ErroLeituraPlanilha("Nesta etapa inicial, selecione uma planilha no formato .xlsx.")

    try:
        workbook = load_workbook(caminho, data_only=True, read_only=False)
    except Exception as erro:
        raise ErroLeituraPlanilha(f"Não foi possível abrir a planilha de origem: {erro}") from erro

    worksheet, linha_cabecalho = _encontrar_planilha_e_linha_cabecalho(workbook)
    mapa_cabecalhos = _mapear_cabecalhos(worksheet, linha_cabecalho)
    colunas = _montar_mapa_colunas(mapa_cabecalhos)
    primeira_linha_dados = linha_cabecalho + 1

    colaboradores: list[dict[str, Any]] = []

    for numero_linha in range(primeira_linha_dados, worksheet.max_row + 1):
        registro = _ler_registro(numero_linha, worksheet, colunas)

        if all(valor_vazio(registro.get(campo)) for campo in colunas):
            continue

        registro["linha_origem"] = numero_linha
        colaboradores.append(registro)

    if not colaboradores:
        raise ErroLeituraPlanilha("Nenhum colaborador foi encontrado na planilha selecionada.")

    return colaboradores
