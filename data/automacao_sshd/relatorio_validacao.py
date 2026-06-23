from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .validacao import PendenciaValidacao


VERDE = "22C55E"
VERMELHO = "EF4444"
FUNDO_CABECALHO = "0F1F17"
TEXTO_CLARO = "F4F7F5"

ABA_EXECUCOES = "Execucoes"
ABA_DETALHES = "Detalhes"

CABECALHO_EXECUCOES = [
    "ID execução",
    "Gerado em",
    "Status",
    "Arquivo fonte",
    "Total de profissionais",
    "Total de pendências",
]

CABECALHO_DETALHES = [
    "ID execução",
    "Gerado em",
    "Status",
    "Linha origem",
    "Nome",
    "CPF",
    "Campo",
    "Mensagem",
]


def _formatar_cabecalho(worksheet) -> None:
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color=TEXTO_CLARO)
        cell.fill = PatternFill("solid", fgColor=FUNDO_CABECALHO)
        cell.alignment = Alignment(horizontal="center")


def _ajustar_larguras(worksheet) -> None:
    for coluna in worksheet.columns:
        largura = 0
        letra = get_column_letter(coluna[0].column)
        for cell in coluna:
            valor = "" if cell.value is None else str(cell.value)
            largura = max(largura, len(valor))
        worksheet.column_dimensions[letra].width = min(max(largura + 2, 12), 60)


def _criar_workbook_relatorio() -> Workbook:
    workbook = Workbook()

    execucoes = workbook.active
    execucoes.title = ABA_EXECUCOES
    execucoes.append(CABECALHO_EXECUCOES)
    _formatar_cabecalho(execucoes)

    detalhes = workbook.create_sheet(ABA_DETALHES)
    detalhes.append(CABECALHO_DETALHES)
    _formatar_cabecalho(detalhes)

    return workbook


def _abrir_ou_criar_relatorio(caminho_relatorio: Path) -> Workbook:
    if not caminho_relatorio.exists():
        return _criar_workbook_relatorio()

    workbook = load_workbook(caminho_relatorio)

    if ABA_EXECUCOES not in workbook.sheetnames:
        execucoes = workbook.create_sheet(ABA_EXECUCOES, 0)
        execucoes.append(CABECALHO_EXECUCOES)
        _formatar_cabecalho(execucoes)

    if ABA_DETALHES not in workbook.sheetnames:
        detalhes = workbook.create_sheet(ABA_DETALHES)
        detalhes.append(CABECALHO_DETALHES)
        _formatar_cabecalho(detalhes)

    return workbook


def _proximo_id_execucao(worksheet) -> int:
    ids: list[int] = []

    for row in worksheet.iter_rows(min_row=2, max_col=1, values_only=True):
        valor = row[0]
        if isinstance(valor, int):
            ids.append(valor)
        elif isinstance(valor, str) and valor.isdigit():
            ids.append(int(valor))

    return max(ids, default=0) + 1


def _formatar_status(cell) -> None:
    cell.fill = PatternFill("solid", fgColor=VERDE if cell.value in {"OK", "APROVADO"} else VERMELHO)
    cell.font = Font(bold=True, color=TEXTO_CLARO)


def gerar_relatorio_validacao(
    *,
    caminho_relatorio: str | Path,
    colaboradores: list[dict[str, Any]],
    pendencias: list[PendenciaValidacao],
    caminho_fonte: str | Path,
) -> Path:
    saida = Path(caminho_relatorio)
    saida.parent.mkdir(parents=True, exist_ok=True)

    workbook = _abrir_ou_criar_relatorio(saida)
    execucoes = workbook[ABA_EXECUCOES]
    detalhes = workbook[ABA_DETALHES]

    id_execucao = _proximo_id_execucao(execucoes)
    gerado_em = datetime.now().strftime("%d/%m/%Y %H:%M:%S")
    status = "APROVADO" if not pendencias else "COM PENDÊNCIAS"

    execucoes.append(
        [
            id_execucao,
            gerado_em,
            status,
            str(caminho_fonte),
            len(colaboradores),
            len(pendencias),
        ]
    )
    _formatar_status(execucoes.cell(execucoes.max_row, 3))

    if pendencias:
        for pendencia in pendencias:
            detalhes.append(
                [
                    id_execucao,
                    gerado_em,
                    "PENDÊNCIA",
                    pendencia.linha_origem,
                    pendencia.nome,
                    pendencia.cpf,
                    pendencia.campo,
                    pendencia.mensagem,
                ]
            )
    else:
        for colaborador in colaboradores:
            detalhes.append(
                [
                    id_execucao,
                    gerado_em,
                    "OK",
                    colaborador.get("linha_origem", ""),
                    colaborador.get("nome_completo", ""),
                    colaborador.get("cpf", ""),
                    "-",
                    "Registro validado com sucesso.",
                ]
            )

    for row in detalhes.iter_rows(min_row=2, max_row=detalhes.max_row):
        _formatar_status(row[2])

    _formatar_cabecalho(execucoes)
    _formatar_cabecalho(detalhes)
    _ajustar_larguras(execucoes)
    _ajustar_larguras(detalhes)

    workbook.save(saida)
    return saida
