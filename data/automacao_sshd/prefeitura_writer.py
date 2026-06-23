from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .config import ABA_TEMPLATE, CELULAS_PREFEITURA, CELULAS_SOLICITANTE, VALORES_FIXOS
from .exceptions import ErroLeituraPlanilha
from .normalizacao import nome_aba_seguro, proximo_caminho_disponivel


def preencher_template_prefeitura(
    *,
    caminho_template: str | Path,
    colaboradores: list[dict[str, Any]],
    nome_solicitante: str,
    sshd_solicitante: str,
    cargo_solicitante: str,
    caminho_saida: str | Path,
) -> Path:
    template = Path(caminho_template)
    saida = proximo_caminho_disponivel(Path(caminho_saida))

    if not template.exists():
        raise ErroLeituraPlanilha(f"Template da Prefeitura nao encontrado: {template}")

    workbook = load_workbook(template)

    if ABA_TEMPLATE not in workbook.sheetnames:
        raise ErroLeituraPlanilha(f"Aba obrigatoria nao encontrada no template: {ABA_TEMPLATE}")

    worksheet_template = workbook[ABA_TEMPLATE]
    nomes_usados = set(workbook.sheetnames)

    for indice, colaborador in enumerate(colaboradores, start=1):
        nome_aba = nome_aba_seguro(str(colaborador.get("nome_completo") or ""), indice, nomes_usados)
        worksheet = workbook.copy_worksheet(worksheet_template)
        worksheet.title = nome_aba

        worksheet[CELULAS_SOLICITANTE["nome"]] = nome_solicitante
        worksheet[CELULAS_SOLICITANTE["sshd"]] = sshd_solicitante
        worksheet[CELULAS_SOLICITANTE["cargo"]] = cargo_solicitante

        for campo, celula in CELULAS_PREFEITURA.items():
            worksheet[celula] = colaborador.get(campo)

        for celula, valor in VALORES_FIXOS.items():
            worksheet[celula] = valor

    workbook.remove(worksheet_template)
    workbook.active = 0
    saida.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(saida)

    return saida
